"""Módulo de exportación y formateo de libros de cálculo Excel (openpyxl).
Provee constructores de encabezados, estilos y utilitarios para los reportes oficiales:
- Manifiesto de Gasto
- FE-05 Ejecución Presupuestal Mensual
- FE-06 Presupuesto vs Ejecutado
- Resumen Financiero (Panel)
"""
import math
import os
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from . import temas as _temas
except (ImportError, ValueError):
    import temas as _temas


def build_xl_header(ws, ncols, title, subtitle, data, logosrc,
                    info_pairs, responsables, extra_right=None,
                    title_font_size=16, logo_w=60, logo_h=60,
                    header_heights=None, info_cols=None, tema=""):
    """Encabezado Excel limpio en 2 secciones (fiel a las hojas de impresión oficial).

    header_heights: dict opcional {row_key: height} para sobreescribir alturas.
        Keys: 'entity'=row1, 'sub'=row2, 'title'=row3, 'subtitle'=row4,
              'project'=row6 info, 'pair'=rows7-8, 'resp'=rows responsables.
    """
    FN = "Calibri"
    PE = _temas.paleta_excel(tema or "")
    AZUL = PE["AZUL"]
    AZUL_OSC = PE["AZUL_OSC"]
    BG = PE["BG"]
    TEXTO2 = PE["TEXTO2"]
    BLANCO = "FFFFFF"
    BORDE = PE["BORDE"]

    thin_o = Side(style="thin", color=AZUL_OSC)
    border_all = Border(left=thin_o, right=thin_o, top=thin_o, bottom=thin_o)
    thin_g = Side(style="thin", color=BORDE)
    border_grid = Border(left=thin_g, right=thin_g, top=thin_g, bottom=thin_g)

    fill_azul = PatternFill("solid", fgColor=AZUL)
    fill_bg = PatternFill("solid", fgColor=BG)
    fill_resp = PatternFill("solid", fgColor=PE["TINTE"])

    font_entity = Font(name=FN, size=11, bold=True, color=AZUL_OSC)
    font_sub = Font(name=FN, size=9, color=TEXTO2)
    font_title = Font(name=FN, size=title_font_size, bold=True, color=AZUL)
    font_label = Font(name=FN, size=8, bold=True, color=AZUL_OSC)
    font_val = Font(name=FN, size=9)
    font_resp_lab = Font(name=FN, size=7, bold=True, color=BLANCO)
    font_resp_val = Font(name=FN, size=9)
    font_extra = Font(name=FN, size=9, bold=True, color=AZUL)

    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_nw = Alignment(horizontal="left", vertical="center", wrap_text=False)
    align_r = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # ── Logo institucional ──
    if logosrc and os.path.isfile(logosrc):
        try:
            img = XlImage(logosrc)
            img.width = logo_w
            img.height = logo_h
            ws.add_image(img, "A1")
        except Exception:
            pass

    pcol = get_column_letter(ncols - 1)
    lcol = get_column_letter(ncols)

    # ── SECCIÓN 1: Encabezado limpio sin bordes alrededor del logo ──
    ws.merge_cells(f"B1:{pcol}1")
    c = ws["B1"]
    c.value = data.get("entidad", "")
    c.font = font_entity
    c.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)

    ws.merge_cells(f"B2:{pcol}2")
    c = ws["B2"]
    c.value = data.get("unidad", "")
    c.font = font_sub
    c.alignment = align_c

    ws.merge_cells(f"B3:{pcol}3")
    c = ws["B3"]
    c.value = title.upper()
    c.font = font_title
    c.alignment = Alignment(horizontal="center", vertical="top")

    ws.merge_cells(f"B4:{pcol}4")
    c = ws["B4"]
    c.value = subtitle.upper()
    c.font = font_sub
    c.alignment = align_c

    # Período / extra a la derecha — sin borde en encabezado
    if extra_right:
        for rn, (val, fnt) in extra_right.items():
            c = ws.cell(row=rn, column=ncols)
            c.value = val
            c.font = fnt or font_extra
            c.alignment = align_r
    elif data.get("periodo"):
        ws.merge_cells(f"{lcol}2:{lcol}3")
        c = ws[f"{lcol}2"]
        c.value = data["periodo"]
        c.font = Font(name=FN, size=9, bold=True, color=AZUL)
        c.alignment = align_r

    # Alturas de encabezado (configurables)
    hh = header_heights or {}
    ws.row_dimensions[1].height = hh.get('entity', 28)
    ws.row_dimensions[2].height = hh.get('sub', 18)
    ws.row_dimensions[3].height = hh.get('title', 32)
    ws.row_dimensions[4].height = hh.get('subtitle', 18)

    # ── SECCIÓN 2: Info-grid con cards bordeadas ──
    info_start = 6

    for i, row_items in enumerate(info_pairs):
        r = info_start + i
        n = len(row_items)
        if info_cols == "left":
            # Labels A:C, values D+ (no wrap) — FE-05
            if n == 1:
                lbl, val = row_items[0]
                ws.merge_cells(f"A{r}:C{r}")
                c_l = ws.cell(row=r, column=1, value=lbl)
                c_l.font = font_label
                c_l.alignment = align_nw
                c_l.fill = fill_bg
                c_l.border = border_grid
                for cc in [2, 3]:
                    ws.cell(row=r, column=cc).fill = fill_bg
                    ws.cell(row=r, column=cc).border = border_grid
                ws.merge_cells(f"D{r}:{lcol}{r}")
                c_v = ws.cell(row=r, column=4, value=val)
                c_v.font = font_val
                c_v.alignment = align_nw
                c_v.border = border_grid
                for cc in range(5, ncols + 1):
                    ws.cell(row=r, column=cc).border = border_grid
            else:
                # Dos mitades: A:C label, D:mid val | mid+1:mid+3 label, mid+4:lastcol val
                mid = (ncols + 1) // 2
                for idx, (lbl, val) in enumerate(row_items):
                    if idx == 0:
                        lbl_s, lbl_e = 1, 3
                        val_s, val_e = 4, mid
                    else:
                        lbl_s, lbl_e = mid + 1, mid + 3
                        val_s, val_e = mid + 4, ncols
                    cl1, cl2 = get_column_letter(lbl_s), get_column_letter(lbl_e)
                    cv1, cv2 = get_column_letter(val_s), get_column_letter(val_e)
                    ws.merge_cells(f"{cl1}{r}:{cl2}{r}")
                    c_l = ws.cell(row=r, column=lbl_s, value=lbl)
                    c_l.font = font_label
                    c_l.alignment = align_nw
                    c_l.fill = fill_bg
                    c_l.border = border_grid
                    for cc in range(lbl_s + 1, lbl_e + 1):
                        ws.cell(row=r, column=cc).fill = fill_bg
                        ws.cell(row=r, column=cc).border = border_grid
                    ws.merge_cells(f"{cv1}{r}:{cv2}{r}")
                    c_v = ws.cell(row=r, column=val_s, value=val)
                    c_v.font = font_val
                    c_v.alignment = align_nw
                    c_v.border = border_grid
                    for cc in range(val_s + 1, val_e + 1):
                        ws.cell(row=r, column=cc).border = border_grid
            ws.row_dimensions[r].height = hh.get('pair', 17)
        elif n == 1:
            # Ancho completo: label A:B, valor C:lastcol
            lbl, val = row_items[0]
            ws.merge_cells(f"A{r}:B{r}")
            c_l = ws.cell(row=r, column=1, value=lbl)
            c_l.font = font_label
            c_l.alignment = align_l
            c_l.fill = fill_bg
            c_l.border = border_grid
            ws.cell(row=r, column=2).fill = fill_bg
            ws.cell(row=r, column=2).border = border_grid
            ws.merge_cells(f"C{r}:{lcol}{r}")
            c_v = ws.cell(row=r, column=3, value=val)
            c_v.font = font_val
            c_v.alignment = align_l
            c_v.border = border_grid
            for cc in range(4, ncols + 1):
                ws.cell(row=r, column=cc).border = border_grid
            v_len = len(str(val or ""))
            nlines = max(1, math.ceil(v_len / 50))
            ws.row_dimensions[r].height = hh.get('project', max(17, nlines * 14))
        else:
            # Dos mitades: par de cards lado a lado
            mid = (ncols + 1) // 2
            for idx, (lbl, val) in enumerate(row_items):
                if idx == 0:
                    col_lbl_s, col_lbl_e = 1, 2
                    col_val_s, col_val_e = 3, mid
                else:
                    col_lbl_s, col_lbl_e = mid + 1, mid + 2
                    col_val_s, col_val_e = mid + 3, ncols
                cl1 = get_column_letter(col_lbl_s)
                cl2 = get_column_letter(col_lbl_e)
                cv1 = get_column_letter(col_val_s)
                cv2 = get_column_letter(col_val_e)
                ws.merge_cells(f"{cl1}{r}:{cl2}{r}")
                c_l = ws.cell(row=r, column=col_lbl_s, value=lbl)
                c_l.font = font_label
                c_l.alignment = align_l
                c_l.fill = fill_bg
                c_l.border = border_grid
                ws.cell(row=r, column=col_lbl_e).fill = fill_bg
                ws.cell(row=r, column=col_lbl_e).border = border_grid
                ws.merge_cells(f"{cv1}{r}:{cv2}{r}")
                c_v = ws.cell(row=r, column=col_val_s, value=val)
                c_v.font = font_val
                c_v.alignment = align_l
                c_v.border = border_grid
                for cc in range(col_val_s + 1, col_val_e + 1):
                    ws.cell(row=r, column=cc).border = border_grid
            ws.row_dimensions[r].height = hh.get('pair', 17)

    # ── Responsables ──
    resp_start = info_start + len(info_pairs)
    resp_col_end = min(5, ncols)
    resp_merge_end = get_column_letter(resp_col_end)
    for i, (lab, val) in enumerate(responsables):
        r = resp_start + i
        if info_cols == "left":
            # Labels A:C, values D:lastcol (sin salto) — FE-05
            ws.merge_cells(f"A{r}:C{r}")
            c_l = ws.cell(row=r, column=1, value=lab)
            c_l.font = font_resp_lab
            c_l.alignment = align_nw
            c_l.fill = fill_azul
            c_l.border = border_all
            for cc in [2, 3]:
                ws.cell(row=r, column=cc).fill = fill_azul
                ws.cell(row=r, column=cc).border = border_all
            ws.merge_cells(f"D{r}:{lcol}{r}")
            c_v = ws.cell(row=r, column=4, value=val)
            c_v.font = font_resp_val
            c_v.alignment = align_nw
            c_v.fill = fill_resp
            c_v.border = border_all
            for cc in range(5, ncols + 1):
                ws.cell(row=r, column=cc).fill = fill_resp
                ws.cell(row=r, column=cc).border = border_all
        else:
            ws.merge_cells(f"A{r}:B{r}")
            c_l = ws.cell(row=r, column=1, value=lab)
            c_l.font = font_resp_lab
            c_l.alignment = align_l
            c_l.fill = fill_azul
            c_l.border = border_all
            ws.cell(row=r, column=2).fill = fill_azul
            ws.cell(row=r, column=2).border = border_all
            ws.merge_cells(f"C{r}:{resp_merge_end}{r}")
            c_v = ws.cell(row=r, column=3, value=val)
            c_v.font = font_resp_val
            c_v.alignment = align_l
            c_v.fill = fill_resp
            c_v.border = border_all
        for cc in range(4, resp_col_end + 1):
            ws.cell(row=r, column=cc).fill = fill_resp
            ws.cell(row=r, column=cc).border = border_all
        for cc in range(resp_col_end + 1, ncols + 1):
            ws.cell(row=r, column=cc).border = border_grid
        ws.row_dimensions[r].height = hh.get('resp', 19)

    return resp_start + len(responsables) + 1
